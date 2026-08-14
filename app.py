import os
import io
import csv
import json
import calendar
from functools import wraps
from datetime import datetime, date, timedelta
from flask import Flask, render_template, request, redirect, url_for, jsonify, flash, Response, session
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash

BASE_DIR = os.path.abspath(os.path.dirname(__file__))

db = SQLAlchemy()

UNITS = ['Sir Q Expense', 'Management', 'Cluster1', 'Cluster2', 'Cluster3',
         'Cluster4', 'Cluster5', 'Cluster6', 'Internal Audit']

ACCOUNT_TITLES = ['Transportation & Travel', 'Meal Allowance', 'Employees Meal Allowance',
                   'Employees Housing Allowance', 'Water', 'Miscellaneous']

# Canonical Excel header titles (as shown in the ledger table) mapped to model fields
IMPORT_HEADERS = {
    'Payee': 'payee',
    'Trxn Code': 'trxn_code',
    'Branch': 'branch',
    'Trxn Date': 'trxn_date',
    'Particulars': 'particulars',
    'Ref#': 'ref_no',
    'Account Title': 'account_title',
    'TIN': 'tin',
    'Address': 'address',
    'Amount': 'amount',
    'Prepared By': 'prepared_by',
}

# Flexible matching: normalized header text -> model field
IMPORT_HEADER_SYNONYMS = {
    'payee': 'payee',
    'branch': 'branch',
    'branchoffice': 'branch',
    'trxndate': 'trxn_date',
    'trxndt': 'trxn_date',
    'transactiondate': 'trxn_date',
    'date': 'trxn_date',
    'dateoftransaction': 'trxn_date',
    'particulars': 'particulars',
    'description': 'particulars',
    'details': 'particulars',
    'ref': 'ref_no',
    'refno': 'ref_no',
    'trxncode': 'trxn_code',
    'transactioncode': 'trxn_code',
    'code': 'trxn_code',
    'reference': 'ref_no',
    'referenceno': 'ref_no',
    'invoiceno': 'ref_no',
    'receiptno': 'ref_no',
    'ornumber': 'ref_no',
    'accounttitle': 'account_title',
    'accttitle': 'account_title',
    'acct': 'account_title',
    'account': 'account_title',
    'category': 'account_title',
    'expenseaccount': 'account_title',
    'tin': 'tin',
    'tinno': 'tin',
    'tinnumber': 'tin',
    'address': 'address',
    'payeeaddress': 'address',
    'amount': 'amount',
    'amountp': 'amount',
    'amountphp': 'amount',
    'amountpeso': 'amount',
    'amountpesos': 'amount',
    'pesoamount': 'amount',
    'peso': 'amount',
    'total': 'amount',
    'totalamount': 'amount',
    'cost': 'amount',
    'preparedby': 'prepared_by',
    'prepared': 'prepared_by',
}

# Short display columns for the dashboard, matching the original sheet's layout
DASHBOARD_COLS = [
    ('Transportation & Travel', 'Transportation & Travel'),
    ('Meal Allowance', 'Meal Allowance'),
    ('Employees Housing Allowance', 'Housing Allowance'),
    ('Water', 'Water'),
    ('Miscellaneous', 'Miscellaneous'),
]


class User(db.Model):
    __tablename__ = 'users'
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(64), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class Expense(db.Model):
    __tablename__ = 'expenses'
    id = db.Column(db.Integer, primary_key=True)
    unit = db.Column(db.String(64), nullable=False, index=True)
    payee = db.Column(db.String(255))
    branch = db.Column(db.String(128))
    trxn_date = db.Column(db.Date, index=True)
    particulars = db.Column(db.Text)
    ref_no = db.Column(db.String(128))
    trxn_code = db.Column(db.String(128))
    account_title = db.Column(db.String(64), index=True)
    tin = db.Column(db.String(64))
    address = db.Column(db.String(255))
    amount = db.Column(db.Float, nullable=False, default=0)
    prepared_by = db.Column(db.String(128), default='Jenny Rose Ando')
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

    def to_dict(self):
        return {
            'id': self.id,
            'unit': self.unit,
            'payee': self.payee,
            'branch': self.branch,
            'trxn_date': self.trxn_date.isoformat() if self.trxn_date else None,
            'particulars': self.particulars,
            'ref_no': self.ref_no,
            'trxn_code': self.trxn_code,
            'account_title': self.account_title,
            'tin': self.tin,
            'address': self.address,
            'amount': self.amount,
            'prepared_by': self.prepared_by,
        }


def parse_date(value):
    if not value:
        return None
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and value > 20000:
        return (datetime(1899, 12, 30) + timedelta(days=value)).date()
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m-%d-%Y', '%Y/%m/%d',
                '%d/%m/%Y', '%d-%m-%Y', '%d %b %Y', '%b %d, %Y',
                '%d-%b-%Y', '%d/%b/%Y', '%b/%d/%Y'):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def _norm_header(text):
    return ''.join(c for c in str(text or '').lower() if c.isalnum()).strip()


def resolve_import_header(text):
    if text is None:
        return None
    key = _norm_header(text)
    if not key:
        return None
    return IMPORT_HEADER_SYNONYMS.get(key)


def _sheet_rows(ws, max_rows=10000):
    """Collect non-empty rows, stopping early after 25 consecutive empty rows."""
    rows = []
    empty_run = 0
    for row in ws.iter_rows(values_only=True):
        if any(v is not None and str(v).strip() != '' for v in row):
            rows.append(list(row))
            empty_run = 0
            if len(rows) >= max_rows:
                break
        else:
            empty_run += 1
            if empty_run >= 25:
                break
    return rows


def _locate_header(rows):
    """Find the row with the most recognized headers among the first 10 rows."""
    best_idx = -1
    best_count = 0
    for i, row in enumerate(rows[:10]):
        count = sum(1 for c in row if resolve_import_header(c) is not None)
        if count > best_count:
            best_count = count
            best_idx = i
    if best_idx < 0 or best_count < 2:
        return [], rows, 0
    return rows[best_idx], rows[best_idx + 1:], best_idx


def workbook_rows(file_storage, filename, max_rows=10000):
    """Return (header_row, data_rows, dropped_title_rows) from the best-matching sheet."""
    if filename.lower().endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_storage.read()), data_only=True, read_only=True)
        best = None
        best_count = -1
        try:
            for ws in wb.worksheets:
                rows = _sheet_rows(ws, max_rows)
                headers, data, dropped = _locate_header(rows)
                count = sum(1 for h in headers if resolve_import_header(h))
                if count > best_count:
                    best_count = count
                    best = (headers, data, dropped)
        finally:
            wb.close()
        return best or ([], [], 0)
    raw = io.BytesIO(file_storage.read()).getvalue()
    text = raw.decode('utf-8-sig', errors='replace')
    rows = [r for r in csv.reader(io.StringIO(text)) if any(c.strip() for c in r)]
    return _locate_header(rows)


def parse_excel_value(value):
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def parse_excel_amount(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(',', '').replace('₱', '').replace('PHP', '')
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def create_app():
    app = Flask(__name__)
    app.config['SECRET_KEY'] = 'revolving-fund-dev-key'
    app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///' + os.path.join(app.instance_path, 'revfund.db')
    app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

    db.init_app(app)

    with app.app_context():
        db.create_all()
        migrate_trxn_code()
        seed_if_empty()
        seed_user_if_empty()

    register_routes(app)
    return app


def seed_user_if_empty():
    if User.query.first():
        return
    db.session.add(User(username='Jenny', password_hash=generate_password_hash('jenny_081226')))
    db.session.commit()


def migrate_trxn_code():
    with db.engine.connect() as conn:
        cols = {r[1] for r in conn.exec_driver_sql('PRAGMA table_info(expenses)')}
        if 'trxn_code' not in cols:
            conn.exec_driver_sql('ALTER TABLE expenses ADD COLUMN trxn_code VARCHAR(128)')
            conn.commit()


def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get('user'):
            flash('Please log in to continue.', 'warning')
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return wrapper


def seed_if_empty():
    if Expense.query.first() is not None:
        return
    seed_path = os.path.join(BASE_DIR, 'seed_data.json')
    if not os.path.exists(seed_path):
        return
    with open(seed_path) as f:
        rows = json.load(f)
    for r in rows:
        db.session.add(Expense(
            unit=r['unit'],
            payee=r['payee'],
            branch=r['branch'],
            trxn_date=parse_date(r['trxn_date']),
            particulars=r['particulars'],
            ref_no=r['ref_no'],
            account_title=r['account_title'],
            tin=r['tin'],
            address=r['address'],
            amount=r['amount'],
        ))
    db.session.commit()


def month_bounds(year, month):
    first = date(year, month, 1)
    last_day = calendar.monthrange(year, month)[1]
    last = date(year, month, last_day)
    return first, last


def build_dashboard(year, month):
    first, last = month_bounds(year, month)

    month_expenses = Expense.query.filter(Expense.trxn_date >= first, Expense.trxn_date <= last).all()

    # Per-unit x account-title matrix for the selected month
    matrix = {u: {a: 0.0 for a, _ in DASHBOARD_COLS} for u in UNITS}
    for e in month_expenses:
        if e.unit in matrix and e.account_title in matrix[e.unit]:
            matrix[e.unit][e.account_title] += e.amount

    unit_rows = []
    grand_totals = {a: 0.0 for a, _ in DASHBOARD_COLS}
    grand_total_all = 0.0
    for u in UNITS:
        row_total = sum(matrix[u].values())
        unit_rows.append({'unit': u, 'amounts': matrix[u], 'total': row_total})
        for a, _ in DASHBOARD_COLS:
            grand_totals[a] += matrix[u][a]
        grand_total_all += row_total

    for row in unit_rows:
        row['pct'] = (row['total'] / grand_total_all * 100) if grand_total_all else 0

    key_metrics = {
        'total_expenses': grand_total_all,
        'num_transactions': len(month_expenses),
        'transportation': grand_totals.get('Transportation & Travel', 0),
        'meal': grand_totals.get('Meal Allowance', 0),
        'housing': grand_totals.get('Employees Housing Allowance', 0),
        'water': grand_totals.get('Water', 0),
        'misc': grand_totals.get('Miscellaneous', 0),
    }

    # Full-year monthly summary (Jan-Dec of the selected year)
    monthly_rows = []
    year_grand = {a: 0.0 for a, _ in DASHBOARD_COLS}
    year_grand_total = 0.0
    for m in range(1, 13):
        mfirst, mlast = month_bounds(year, m)
        rows = Expense.query.filter(Expense.trxn_date >= mfirst, Expense.trxn_date <= mlast).all()
        totals = {a: 0.0 for a, _ in DASHBOARD_COLS}
        for e in rows:
            if e.account_title in totals:
                totals[e.account_title] += e.amount
        total = sum(totals.values())
        monthly_rows.append({
            'label': calendar.month_name[m] + f' {year}',
            'month': m,
            'amounts': totals,
            'total': total,
        })
        for a, _ in DASHBOARD_COLS:
            year_grand[a] += totals[a]
        year_grand_total += total

    for row in monthly_rows:
        row['pct'] = (row['total'] / year_grand_total * 100) if year_grand_total else 0

    return {
        'year': year,
        'month': month,
        'month_label': calendar.month_name[month],
        'key_metrics': key_metrics,
        'unit_rows': unit_rows,
        'grand_totals': grand_totals,
        'grand_total_all': grand_total_all,
        'monthly_rows': monthly_rows,
        'year_grand': year_grand,
        'year_grand_total': year_grand_total,
    }


def register_routes(app):

    @app.context_processor
    def inject_current_user():
        return {'current_user': session.get('user')}

    @app.route('/login', methods=['GET', 'POST'])
    def login():
        if session.get('user'):
            return redirect(url_for('dashboard'))
        if request.method == 'POST':
            username = request.form.get('username', '').strip()
            password = request.form.get('password', '')
            user = User.query.filter_by(username=username).first()
            if user and check_password_hash(user.password_hash, password):
                session['user'] = user.username
                flash(f'Welcome back, {user.username}.', 'success')
                return redirect(url_for('dashboard'))
            flash('Invalid username or password.', 'danger')
        return render_template('login.html')

    @app.route('/logout')
    def logout():
        session.pop('user', None)
        flash('You have been logged out.', 'info')
        return redirect(url_for('login'))

    @app.route('/')
    def index():
        return redirect(url_for('dashboard'))

    @app.route('/dashboard')
    @login_required
    def dashboard():
        if 'month' not in request.args or 'year' not in request.args:
            latest = Expense.query.order_by(Expense.trxn_date.desc().nullslast()).first()
            if latest and latest.trxn_date:
                today = latest.trxn_date
            else:
                today = date.today()
        else:
            today = date.today()
        year = request.args.get('year', type=int, default=today.year)
        month = request.args.get('month', type=int, default=today.month)
        data = build_dashboard(year, month)
        return render_template('dashboard.html', cols=DASHBOARD_COLS, units=UNITS, **data)

    @app.route('/unit/<unit_name>')
    @login_required
    def unit_ledger(unit_name):
        if unit_name not in UNITS:
            flash('Unknown unit', 'danger')
            return redirect(url_for('dashboard'))
        q = Expense.query.filter_by(unit=unit_name)
        account_filter = request.args.get('account_title')
        if account_filter:
            q = q.filter_by(account_title=account_filter)
        entries = q.order_by(Expense.trxn_date.asc().nullslast(), Expense.id.asc()).all()
        total = sum(e.amount for e in entries)
        last_import = session.get('last_import')
        return render_template('ledger.html', unit=unit_name, units=UNITS, entries=entries,
                                total=total, account_titles=ACCOUNT_TITLES,
                                account_filter=account_filter,
                                last_import=last_import if last_import and last_import.get('unit') == unit_name else None)

    @app.route('/unit/<unit_name>/add', methods=['POST'])
    @login_required
    def add_entry(unit_name):
        if unit_name not in UNITS:
            return jsonify({'error': 'Unknown unit'}), 400
        form = request.form
        try:
            amount = float(form.get('amount', 0) or 0)
        except ValueError:
            amount = 0
        entry = Expense(
            unit=unit_name,
            payee=form.get('payee', '').strip(),
            branch=form.get('branch', '').strip() or None,
            trxn_date=parse_date(form.get('trxn_date')),
            particulars=form.get('particulars', '').strip() or None,
            ref_no=form.get('ref_no', '').strip() or None,
            trxn_code=form.get('trxn_code', '').strip() or None,
            account_title=form.get('account_title') or None,
            tin=form.get('tin', '').strip() or None,
            address=form.get('address', '').strip() or None,
            amount=amount,
            prepared_by=form.get('prepared_by', 'Jenny Rose Ando').strip() or 'Jenny Rose Ando',
        )
        db.session.add(entry)
        db.session.commit()
        flash('Expense entry added.', 'success')
        return redirect(url_for('unit_ledger', unit_name=unit_name))

    @app.route('/unit/<unit_name>/import', methods=['POST'])
    @login_required
    def import_entries(unit_name):
        if unit_name not in UNITS:
            flash('Unknown unit', 'danger')
            return redirect(url_for('dashboard'))
        file_storage = request.files.get('file')
        if not file_storage or not file_storage.filename:
            flash('No file selected.', 'warning')
            return redirect(url_for('unit_ledger', unit_name=unit_name))
        filename = file_storage.filename.lower()
        if not (filename.endswith('.xlsx') or filename.endswith('.xlsm') or filename.endswith('.csv')):
            flash('Unsupported file type. Please upload an .xlsx or .csv file.', 'warning')
            return redirect(url_for('unit_ledger', unit_name=unit_name))
        try:
            headers, rows, title_dropped = workbook_rows(file_storage, filename)
        except Exception:
            flash('Could not read the file. Please check that it is a valid Excel/CSV file.', 'danger')
            return redirect(url_for('unit_ledger', unit_name=unit_name))
        if len(rows) < 1:
            flash('File has no data rows (header row only or empty).', 'warning')
            return redirect(url_for('unit_ledger', unit_name=unit_name))

        col_index = {}
        unknown = []
        for i, h in enumerate(headers):
            field = resolve_import_header(h)
            if field:
                col_index[i] = field
            else:
                unknown.append(h)
        if 'payee' not in col_index.values() or 'amount' not in col_index.values():
            flash('Required columns "Payee" and "Amount" were not found. '
                  'Expected headers (case/format-flexible): ' + ', '.join(IMPORT_HEADERS), 'danger')
            return redirect(url_for('unit_ledger', unit_name=unit_name))

        imported = 0
        skipped_incomplete = 0
        skipped_bad_title = 0
        skipped_total = 0
        import_boundary = db.session.query(db.func.max(Expense.id)).scalar() or 0
        for row in rows:
            data = {}
            for i, field in col_index.items():
                data[field] = row[i] if i < len(row) else None
            payee = parse_excel_value(data.get('payee'))
            amount = parse_excel_amount(data.get('amount'))
            if payee and _norm_header(payee).startswith('total'):
                skipped_total += 1
                continue
            if not payee or amount is None:
                skipped_incomplete += 1
                continue
            account_title = parse_excel_value(data.get('account_title'))
            if account_title and account_title not in ACCOUNT_TITLES:
                skipped_bad_title += 1
                continue
            db.session.add(Expense(
                unit=unit_name,
                payee=payee,
                branch=parse_excel_value(data.get('branch')),
                trxn_date=parse_date(data.get('trxn_date')),
                particulars=parse_excel_value(data.get('particulars')),
                ref_no=parse_excel_value(data.get('ref_no')),
                trxn_code=parse_excel_value(data.get('trxn_code')),
                account_title=account_title,
                tin=parse_excel_value(data.get('tin')),
                address=parse_excel_value(data.get('address')),
                amount=amount,
                prepared_by=parse_excel_value(data.get('prepared_by')) or 'Jenny Rose Ando',
            ))
            imported += 1
        db.session.commit()
        capped = ' (stopped at 10,000 rows)' if len(rows) >= 10000 else ''
        message = f'Import complete: {imported} row(s) added{capped}.'
        if imported:
            session['last_import'] = {'unit': unit_name, 'min_id': import_boundary + 1, 'count': imported}
        flash(message, 'success' if imported else 'warning')
        return redirect(url_for('unit_ledger', unit_name=unit_name))

    @app.route('/unit/<unit_name>/delete-import', methods=['POST'])
    @login_required
    def delete_import(unit_name):
        if unit_name not in UNITS:
            return jsonify({'error': 'Unknown unit'}), 400
        last = session.get('last_import')
        if not last or last.get('unit') != unit_name:
            flash('No recent import to delete for this unit.', 'warning')
            return redirect(url_for('unit_ledger', unit_name=unit_name))
        deleted = Expense.query.filter(
            Expense.unit == unit_name, Expense.id >= last['min_id']
        ).delete(synchronize_session=False)
        db.session.commit()
        session.pop('last_import', None)
        flash(f'Deleted {deleted} row(s) from the last import.', 'info')
        return redirect(url_for('unit_ledger', unit_name=unit_name))

    @app.route('/entry/<int:entry_id>/edit', methods=['POST'])
    @login_required
    def edit_entry(entry_id):
        entry = Expense.query.get_or_404(entry_id)
        form = request.form
        try:
            amount = float(form.get('amount', entry.amount) or 0)
        except ValueError:
            amount = entry.amount
        entry.payee = form.get('payee', entry.payee).strip()
        entry.branch = form.get('branch', entry.branch)
        entry.trxn_date = parse_date(form.get('trxn_date')) or entry.trxn_date
        entry.particulars = form.get('particulars', entry.particulars)
        entry.ref_no = form.get('ref_no', entry.ref_no)
        entry.trxn_code = form.get('trxn_code', entry.trxn_code).strip() or None
        entry.account_title = form.get('account_title', entry.account_title)
        entry.tin = form.get('tin', entry.tin)
        entry.address = form.get('address', entry.address)
        entry.amount = amount
        entry.prepared_by = form.get('prepared_by', entry.prepared_by).strip() or 'Jenny Rose Ando'
        db.session.commit()
        flash('Expense entry updated.', 'success')
        return redirect(url_for('unit_ledger', unit_name=entry.unit))

    @app.route('/entry/<int:entry_id>/delete', methods=['POST'])
    @login_required
    def delete_entry(entry_id):
        entry = Expense.query.get_or_404(entry_id)
        unit_name = entry.unit
        db.session.delete(entry)
        db.session.commit()
        flash('Expense entry deleted.', 'info')
        return redirect(url_for('unit_ledger', unit_name=unit_name))

    @app.route('/api/expenses')
    @login_required
    def api_expenses():
        unit = request.args.get('unit')
        q = Expense.query
        if unit:
            q = q.filter_by(unit=unit)
        return jsonify([e.to_dict() for e in q.all()])

    @app.route('/export/<unit_name>.csv')
    @login_required
    def export_csv(unit_name):
        if unit_name not in UNITS:
            return 'Unknown unit', 404
        entries = Expense.query.filter_by(unit=unit_name).order_by(Expense.trxn_date.asc()).all()
        lines = ['Payee,Trxn Code,Branch,Transaction Date,Particulars,Ref#,Account Title,TIN,Address,Amount']
        for e in entries:
            row = [e.payee, e.trxn_code, e.branch, e.trxn_date.isoformat() if e.trxn_date else '',
                   e.particulars, e.ref_no, e.account_title, e.tin, e.address, str(e.amount)]
            row = ['' if v is None else str(v).replace(',', ';').replace('\n', ' ') for v in row]
            lines.append(','.join(row))
        csv_data = '\n'.join(lines)
        return Response(csv_data, mimetype='text/csv',
                         headers={'Content-Disposition': f'attachment; filename={unit_name}_expenses.csv'})


app = create_app()

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
